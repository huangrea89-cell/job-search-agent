from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.db import Database
from app.models import SourceJob, UserStatus
from app.resume import ResumeService, detect_language
from app.scoring import CandidateEvidence
from app.service import JobService

ROOT = Path(__file__).parents[1]


def setup_case(tmp_path):
    master = tmp_path / "master.docx"
    document = Document()
    document.add_heading("候选人简历", level=1)
    document.add_paragraph("基于访谈完成用户需求分析。")
    document.save(master)
    db = Database(tmp_path / "db.sqlite3")
    db.migrate()
    jobs = JobService(db, ROOT / "config/scoring.yaml")
    result = jobs.import_job(
        SourceJob(source="official",source_url="https://example.com/1",company="示例",title="AI 产品经理",description="负责大模型产品和用户研究",locations=["上海"],employment_type="full_time"),
        CandidateEvidence(skills=["用户研究"]),
    )
    jobs.set_user_status(result["job_id"], UserStatus.INTERESTED)
    resumes = ResumeService(db)
    version_id = resumes.register_version("v-test", master, "zh")
    change = {
        "location": "经历第1条", "before": "基于访谈完成用户需求分析。",
        "after": "通过访谈完成用户需求分析并形成产品建议。", "job_requirement": "用户研究",
        "evidence_refs": ["fact-1"], "reason": "突出与岗位直接相关的真实证据",
    }
    return resumes, result["job_id"], version_id, master, change


def test_language_detection_stops_on_mixed_content():
    assert detect_language("负责用户研究和产品设计") == "zh"
    assert detect_language("Own product discovery and user research") == "en"
    assert detect_language("负责 AI product discovery") is None


def test_proposal_requires_confirmed_fact_reference(tmp_path):
    resumes, job_id, version_id, _, change = setup_case(tmp_path)
    with pytest.raises(ValueError, match="confirmed"):
        resumes.create_proposal(job_id, version_id, [change], {})


def test_generation_requires_confirmation_and_preserves_master(tmp_path):
    resumes, job_id, version_id, master, change = setup_case(tmp_path)
    original = master.read_bytes()
    proposal_id = resumes.create_proposal(job_id, version_id, [change], {"fact-1": "完成用户访谈并形成产品建议"})
    with pytest.raises(PermissionError):
        resumes.generate_confirmed_docx(proposal_id, [0], tmp_path / "tailored.docx", confirmed=False)
    output = resumes.generate_confirmed_docx(proposal_id, [0], tmp_path / "tailored.docx", confirmed=True)
    assert master.read_bytes() == original


def test_generation_tolerates_layout_only_trailing_spaces(tmp_path):
    resumes, job_id, version_id, master, change = setup_case(tmp_path)
    document = Document(master)
    document.paragraphs[1].text += "   "
    document.save(master)
    version_id = resumes.register_version("v-spaces", master, "zh")
    proposal_id = resumes.create_proposal(job_id, version_id, [change], {"fact-1": "已确认事实"})
    output = resumes.generate_confirmed_docx(proposal_id, [0], tmp_path / "tailored.docx", confirmed=True)
    assert Document(output).paragraphs[1].text == change["after"]
    assert "形成产品建议" in "\n".join(p.text for p in Document(output).paragraphs)
    fonts = list(Document(output).styles.element.iter(qn("w:rFonts")))
    assert any(item.get(qn("w:eastAsia")) == "Arial Unicode MS" for item in fonts)


def test_registered_version_is_immutable(tmp_path):
    resumes, _, _, master, _ = setup_case(tmp_path)
    Document().save(master)
    with pytest.raises(ValueError, match="immutable"):
        resumes.register_version("v-test", master, "zh")


def test_proposal_rejects_multiple_interested_jobs(tmp_path):
    resumes, job_id, version_id, _, change = setup_case(tmp_path)
    service = JobService(resumes.database, ROOT / "config/scoring.yaml")
    second = service.import_job(
        SourceJob(source="official",source_url="https://example.com/2",company="示例2",title="AI 产品运营",description="负责AI产品运营",locations=["上海"],employment_type="full_time"),
        CandidateEvidence(skills=["AI产品"]),
    )["job_id"]
    service.set_user_status(second, UserStatus.INTERESTED)
    with pytest.raises(ValueError, match="exactly one"):
        resumes.create_proposal(job_id, version_id, [change], {"fact-1": "已确认事实"})
