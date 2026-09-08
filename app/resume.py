from __future__ import annotations

import hashlib
import json
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .db import Database
from .models import utc_now


def _iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _set_cjk_font_fallback(document, font_name: str = "Arial Unicode MS") -> None:
    """Use one Unicode font mapping so CJK text survives cross-suite rendering."""
    roots = [document.element, document.styles.element]
    for section in document.sections:
        roots.extend([section.header._element, section.footer._element])
    seen: set[int] = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for fonts in root.iter(qn("w:rFonts")):
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), font_name)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def detect_language(text: str) -> str | None:
    chinese = sum("\u4e00" <= char <= "\u9fff" for char in text)
    latin = sum(char.isascii() and char.isalpha() for char in text)
    if chinese >= 4 and chinese > latin * 0.25:
        return "zh"
    if chinese == 0 or (chinese <= 1 and latin >= 20):
        return "en"
    return None


class ResumeService:
    def __init__(self, database: Database):
        self.database = database

    def register_version(self, version: str, source_path: Path, language: str) -> int:
        if language not in {"zh", "en"}:
            raise ValueError("resume language must be zh or en")
        if source_path.suffix.lower() != ".docx" or not source_path.is_file():
            raise ValueError("master resume must be an existing DOCX")
        digest = sha256_file(source_path)
        with self.database.connect() as connection:
            existing = connection.execute("SELECT id,source_sha256 FROM resume_versions WHERE version=?", (version,)).fetchone()
            if existing:
                if existing["source_sha256"] != digest:
                    raise ValueError("resume version is immutable; use a new version name")
                return int(existing["id"])
            cursor = connection.execute(
                "INSERT INTO resume_versions(version,source_path,source_sha256,language,created_at) VALUES (?,?,?,?,?)",
                (version, str(source_path), digest, language, utc_now()),
            )
        return int(cursor.lastrowid)

    def create_proposal(
        self,
        job_id: int,
        resume_version_id: int,
        changes: list[dict],
        confirmed_facts: dict[str, str],
        language: str | None = None,
    ) -> int:
        with self.database.connect() as connection:
            job = connection.execute("SELECT description,user_status FROM jobs WHERE id=?", (job_id,)).fetchone()
            resume = connection.execute("SELECT language FROM resume_versions WHERE id=?", (resume_version_id,)).fetchone()
            if not job or not resume:
                raise KeyError("job or resume version not found")
            if job["user_status"] != "interested":
                raise ValueError("exactly one job must first be marked interested")
            interested_count = connection.execute(
                "SELECT COUNT(*) FROM jobs WHERE user_status='interested'"
            ).fetchone()[0]
            if interested_count != 1:
                raise ValueError("exactly one job must first be marked interested")
            detected = detect_language(job["description"])
            selected_language = language or detected
            if selected_language is None:
                raise ValueError("mixed-language job requires explicit resume language")
            if selected_language != resume["language"]:
                raise ValueError("selected resume version language does not match the job")
            normalized = []
            for change in changes:
                required = {"location", "before", "after", "job_requirement", "evidence_refs", "reason"}
                if set(change) != required:
                    raise ValueError("proposal change has missing or unexpected fields")
                refs = change["evidence_refs"]
                if not refs or any(ref not in confirmed_facts for ref in refs):
                    raise ValueError("every proposed change must cite confirmed candidate facts")
                normalized.append({**change, "fact_evidence": [confirmed_facts[ref] for ref in refs]})
            payload = {"language": selected_language, "changes": normalized}
            cursor = connection.execute(
                """INSERT INTO resume_proposals(job_id,resume_version_id,status,proposal_json,created_at)
                VALUES (?,?,?,?,?)""",
                (job_id, resume_version_id, "awaiting_confirmation", json.dumps(payload, ensure_ascii=False), utc_now()),
            )
        return int(cursor.lastrowid)

    def generate_confirmed_docx(
        self, proposal_id: int, approved_change_indexes: list[int], output_path: Path, *, confirmed: bool
    ) -> Path:
        if not confirmed:
            raise PermissionError("explicit resume-change confirmation is required")
        with self.database.connect() as connection:
            row = connection.execute(
                """SELECT p.status,p.proposal_json,r.source_path,r.source_sha256
                FROM resume_proposals p JOIN resume_versions r ON r.id=p.resume_version_id WHERE p.id=?""",
                (proposal_id,),
            ).fetchone()
            if not row:
                raise KeyError(proposal_id)
            if row["status"] not in {"awaiting_confirmation", "approved"}:
                raise ValueError("proposal is not eligible for generation")
            source_path = Path(row["source_path"])
            if sha256_file(source_path) != row["source_sha256"]:
                raise ValueError("master resume changed; register a new immutable version")
            payload = json.loads(row["proposal_json"])
        selected = [payload["changes"][index] for index in approved_change_indexes]
        document = Document(source_path)
        for change in selected:
            replaced = False
            for paragraph in _iter_paragraphs(document):
                # Word/WPS may leave layout-only trailing spaces in a paragraph.
                # Ignore those spaces, but keep the substantive text match exact.
                if paragraph.text.rstrip() == change["before"].rstrip():
                    for run in paragraph.runs:
                        run.text = ""
                    if paragraph.runs:
                        paragraph.runs[0].text = change["after"]
                    else:
                        paragraph.add_run(change["after"])
                    replaced = True
                    break
            if not replaced:
                raise ValueError(f"source text not found at {change['location']}")
        _set_cjk_font_fallback(document)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_suffix(".tmp.docx")
        document.save(temporary)
        temporary.replace(output_path)
        with self.database.connect() as connection:
            connection.execute(
                "UPDATE resume_proposals SET status='generated',output_path=?,confirmed_at=? WHERE id=?",
                (str(output_path), utc_now(), proposal_id),
            )
            connection.execute(
                "INSERT INTO audit_log(event_type,entity_type,entity_id,details_json,created_at) VALUES (?,?,?,?,?)",
                ("resume_generated", "resume_proposal", str(proposal_id), json.dumps({"approved_changes": approved_change_indexes}), utc_now()),
            )
        return output_path


def render_docx_to_pdf(docx_path: Path, output_dir: Path, renderer_path: Path, python_path: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [str(python_path), str(renderer_path), str(docx_path), "--output_dir", str(output_dir), "--emit_pdf"],
        check=False,
        capture_output=True,
        text=True,
        env={**os.environ, "TMPDIR": "/private/tmp"},
    )
    pdf = output_dir / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not pdf.is_file() or pdf.stat().st_size == 0:
        raise RuntimeError("DOCX rendering failed; final PDF was not produced")
    return pdf
