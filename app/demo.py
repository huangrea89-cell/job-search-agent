"""Offline portfolio example. Never opens the user's configured database."""
from __future__ import annotations

import json
import tempfile
from html import escape
from pathlib import Path

from .db import Database
from .models import SourceJob, UserStatus
from .resume import ResumeService
from .scoring import CandidateEvidence
from .service import JobService
from .workspace import build_workspace_snapshot
from docx import Document


def run_demo(root: Path, output_parent: Path) -> Path:
    # A new directory on every run prevents overwriting real or prior demo files.
    output_parent.mkdir(parents=True, exist_ok=True)
    output = Path(tempfile.mkdtemp(prefix="demo-", dir=output_parent)).resolve()
    fixture = json.loads((root / "examples/demo.json").read_text(encoding="utf-8"))
    database = Database(output / "demo.sqlite3")
    database.migrate()
    service = JobService(database, root / "config/scoring.yaml")
    evidence = CandidateEvidence(**fixture["candidate"]["evidence"])
    results = [service.import_job(SourceJob(**job), evidence) for job in fixture["jobs"]]
    # Replay a source record to demonstrate idempotent import with the real pipeline.
    replay = service.import_job(SourceJob(**fixture["jobs"][0]), evidence)
    job_id = results[0]["job_id"]
    service.set_user_status(job_id, UserStatus.INTERESTED)
    master = output / "fictional-resume.docx"
    document = Document()
    document.add_heading("演示候选人（完全虚构）", 0)
    document.add_paragraph(fixture["candidate"]["fact"])
    document.save(master)
    resumes = ResumeService(database)
    version = resumes.register_version("fictional-v1", master, "zh")
    proposal_id = resumes.create_proposal(
        job_id, version, [fixture["proposal"]],
        {"demo-fact-1": fixture["candidate"]["fact"]}, language="zh",
    )
    with database.connect() as connection:
        jobs = [dict(row) for row in connection.execute(
            "SELECT j.company,j.title,s.* FROM jobs j JOIN job_scores s ON j.id=s.job_id ORDER BY j.id"
        )]
        proposal = dict(connection.execute(
            "SELECT id,status,proposal_json FROM resume_proposals WHERE id=?", (proposal_id,)
        ).fetchone())
    payload = {"fictional": True, "input_records": len(results) + 1,
               "unique_jobs": len(jobs), "duplicate_created": replay["created"],
               "jobs": jobs, "proposal": proposal}
    (output / "results.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    snapshot_path = build_workspace_snapshot(database, output / "workspace.snapshot.json", "")
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    snapshot.update(demo=True, generated_from="demo.sqlite3")
    snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    report = ["# 离线演示结果", "", "全部为虚构资料；使用真实规则引擎计算，无模型调用、无网络请求。", "",
              f"导入 {len(results)+1} 次 → {len(jobs)} 个独立岗位；重复导入创建新岗位：{replay['created']}。", "",
              "| 岗位 | 期望分 | 履历分 | 总分 | 星级 | 结果 |", "| --- | ---: | ---: | ---: | ---: | --- |"]
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1100" height="560" viewBox="0 0 1100 560">',
           '<rect width="1100" height="560" fill="#f4f7fb"/>',
           '<g font-family="Arial, sans-serif" fill="#16324f">',
           '<text x="40" y="55" font-size="28" font-weight="bold">Job Agent / Offline demo</text>',
           '<text x="40" y="88" font-size="16">Synthetic data · real scoring pipeline · no API key required</text>',
           '<text x="40" y="126" font-size="16">6 imports / 5 jobs / duplicate merged / proposal awaits approval</text>']
    for i, job in enumerate(jobs):
        notes = json.loads(job["reasons_json"]) + json.loads(job["needs_confirmation_json"])
        report.append(f"| {job['company']} / {job['title']} | {job['expectation_score']} | {job['resume_score']} | {job['total_score']} | {job['stars']} | {'过滤' if job['filtered'] else '保留'}；{'；'.join(notes) or '无额外风险'} |")
        y = 154 + 65 * i
        color = "#fee2e2" if job["filtered"] else "#e0f2fe"
        svg.extend([f'<rect x="40" y="{y}" width="1020" height="54" rx="8" fill="{color}"/>',
                    f'<text x="58" y="{y+33}" font-size="18">{escape(fixture["labels"][i])}</text>',
                    f'<text x="655" y="{y+33}" font-size="17">{job["expectation_score"]} + {job["resume_score"]} = {job["total_score"]} / {job["stars"]} stars / {"filtered" if job["filtered"] else "retained"}</text>'])
    svg.extend(['<text x="40" y="525" font-size="15">Generated result preview, not an Excel screenshot. Stars are not hiring probabilities.</text>', '</g></svg>'])
    (output / "preview.svg").write_text("\n".join(svg), encoding="utf-8")
    change = fixture["proposal"]
    report.extend(["", "## 简历方案（预置示例，非实时模型生成）", "", f"原文：{change['before']}", "", f"拟改为：{change['after']}", "", f"原因：{change['reason']}", "", "状态：awaiting_confirmation。演示没有模拟用户批准，也不生成最终定制简历。", "", "运行 demo 命令生成的独立目录包含演示 SQLite、机器可读结果、虚构主简历、工作台快照及结果预览。"])
    (output / "README.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    return output
